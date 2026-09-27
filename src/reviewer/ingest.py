"""Turn an uploaded file into plain text plus a page map.

Everything downstream works on character offsets into one text string. The page map
lets the UI say "page 7" and lets a lawyer open the PDF to the right place. Offsets are
the source of truth; pages are a convenience derived from them.

Supported: PDF (text layer), DOCX, TXT. Scanned PDFs with no text layer are detected and
reported, not silently returned as empty. OCR is a Phase 2b item (see scope.md risks).
"""

from __future__ import annotations

import bisect
import io
from dataclasses import dataclass, field
from pathlib import Path


class IngestError(ValueError):
    pass


@dataclass
class Document:
    text: str
    filename: str
    media_type: str
    page_starts: list[int] = field(default_factory=list)  # char offset where each page begins

    @property
    def pages(self) -> int:
        return max(1, len(self.page_starts))

    def page_of(self, offset: int) -> int:
        """1-based page number containing a character offset."""
        if not self.page_starts:
            return 1
        return bisect.bisect_right(self.page_starts, offset)


def _from_pdf(data: bytes, filename: str) -> Document:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    starts: list[int] = []
    pos = 0
    for page in reader.pages:
        text = page.extract_text() or ""
        starts.append(pos)
        parts.append(text)
        pos += len(text) + 1  # newline between pages
    full = "\n".join(parts)
    if len(full.strip()) < 50 and len(reader.pages) > 0:
        raise IngestError(
            f"{filename}: no usable text layer on {len(reader.pages)} page(s). "
            "Likely a scanned document; OCR is not enabled in this build."
        )
    return Document(text=full, filename=filename, media_type="application/pdf", page_starts=starts)


def _from_docx(data: bytes, filename: str) -> Document:
    import docx

    document = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs)
    return Document(
        text=text,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        page_starts=[0],
    )


def _from_text(data: bytes, filename: str) -> Document:
    text = data.decode("utf-8", errors="replace")
    return Document(text=text, filename=filename, media_type="text/plain", page_starts=[0])


def ingest_bytes(data: bytes, filename: str) -> Document:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _from_pdf(data, filename)
    if suffix == ".docx":
        return _from_docx(data, filename)
    if suffix in {".txt", ".md", ""}:
        return _from_text(data, filename)
    raise IngestError(f"{filename}: unsupported file type {suffix!r}")


def ingest_path(path: Path) -> Document:
    return ingest_bytes(path.read_bytes(), path.name)
